"""Build docs/datasets/Sagar_Setu_Datasets.xlsx: a catalogue sheet plus one sheet per
exported CSV sample. Run scripts/export_datasets.py first. No formulas — every cell is a
value taken from the exported data, so the workbook opens identically everywhere."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from marine.config import ROOT  # noqa: E402

D = ROOT / "docs" / "datasets"
OUT = D / "Sagar_Setu_Datasets.xlsx"
FONT = "Arial"

CATALOGUE = [
    # sheet, file, dataset, publisher, what, role, cadence / resolution, access, status
    (
        "PFZ lines",
        "incois_pfz_lines.csv",
        "Potential Fishing Zone advisory",
        "INCOIS (MoES, Govt of India)",
        "Daily lines where fish are likely, per coastal sector; geometry + state + advisory date",
        "Where to fish: bearing and distance from the harbour",
        "Daily; line geometry",
        "GeoServer WFS (public, undocumented; what INCOIS's own web map calls)",
        "Live, verified 15–16 Sep 2026",
    ),
    (
        "PFZ vertices",
        "incois_pfz_vertices.csv",
        "PFZ advisory — every vertex",
        "INCOIS",
        "The same 87 lines expanded to 42,740 lat/lon points for plotting",
        "Map overlay",
        "Daily",
        "As above",
        "Live",
    ),
    (
        "Open-Meteo marine",
        "open_meteo_marine_hourly.csv",
        "Marine forecast",
        "Open-Meteo (MFWAM / ECMWF wave models)",
        "Wave height, period, direction; swell; wind-waves; SST; currents — hourly, 7 days back + 8 ahead",
        "Safety verdict inputs; the reliability floor of every chain",
        "Hourly; ~5 km grid",
        "REST, no key, CC-BY 4.0",
        "Live, verified",
    ),
    (
        "Open-Meteo marine units",
        "open_meteo_marine_units.csv",
        "Units for the marine sheet",
        "Open-Meteo",
        "",
        "",
        "",
        "",
        "",
    ),
    (
        "Open-Meteo weather",
        "open_meteo_weather_hourly.csv",
        "Weather forecast",
        "Open-Meteo (ECMWF / ICON etc.)",
        "Wind speed (kn), gusts, direction, visibility, precipitation, WMO weather code, CAPE",
        "Safety verdict inputs; convective proxy while no lightning feed",
        "Hourly; 1–10 km",
        "REST, no key",
        "Live, verified",
    ),
    (
        "Open-Meteo weather units",
        "open_meteo_weather_units.csv",
        "Units for the weather sheet",
        "Open-Meteo",
        "",
        "",
        "",
        "",
        "",
    ),
    (
        "Copernicus chl forecast",
        "copernicus_chlorophyll_forecast.csv",
        "Chlorophyll-a forecast (BGC)",
        "Copernicus Marine Service (EU) / Mercator Ocean",
        "Daily chlorophyll-a, 0.25° mean around each point, 1 Jul → 25 Sep 2026 (9-day forecast)",
        "Productivity; anomaly vs baseline",
        "Daily; 0.25°",
        "Python toolbox, free account",
        "Live, verified 16 Sep 2026",
    ),
    (
        "Copernicus chl Kochi 5y",
        "copernicus_chlorophyll_kochi_2021_2026.csv",
        "Chlorophyll-a five-year daily record at Kochi",
        "Copernicus Marine Service",
        "1,792 daily values, Nov 2021 → Sep 2026 — the multi-year baseline for anomaly detection",
        "Baseline for 'why is the catch declining?'",
        "Daily",
        "As above",
        "Live",
    ),
    (
        "NOAA VIIRS chl",
        "noaa_viirs_chlorophyll.csv",
        "Chlorophyll-a satellite (VIIRS, gap-filled)",
        "NOAA CoastWatch (USA)",
        "Daily ~4 km chlorophyll, 15 Aug → 12 Sep 2026 at six points",
        "No-account fallback behind Copernicus; finer coastal detail",
        "Daily; ~4 km; ~3 days latency",
        "ERDDAP REST, no key",
        "Live, verified",
    ),
    (
        "Boundaries",
        "marine_regions_boundaries.csv",
        "Maritime boundaries incl. India–Sri Lanka",
        "Marine Regions (Flanders Marine Institute)",
        "32 line features; 23 segments are true international boundaries used for geofencing",
        "Geofencing, predictive crossing warning",
        "Static",
        "WFS download, CC-BY",
        "Static, committed",
    ),
    (
        "EEZ polygon",
        "marine_regions_eez_polygon.csv",
        "Indian Exclusive Economic Zone",
        "Marine Regions",
        "Polygon metadata (geometry in data/static/)",
        "Inside / outside Indian waters",
        "Static",
        "WFS download",
        "Static, committed",
    ),
    (
        "Landing centres",
        "incois_landing_centres.csv",
        "Fish landing centres",
        "INCOIS",
        "1,223 named centres with district, sector, lat/lon",
        "Place-name → coordinates (no geocoding API)",
        "Static (Apr 2024)",
        "GeoServer WFS",
        "Static, committed",
    ),
    (
        "PFZ sectors",
        "incois_pfz_sectors.csv",
        "PFZ coastal sectors",
        "INCOIS",
        "14 sector polygons (bounding boxes here)",
        "Sector lookup",
        "Static",
        "GeoServer WFS",
        "Static, committed",
    ),
    (
        "Argo temperature",
        "incois_argo_temperature.csv",
        "Argo gridded temperature at 5 m",
        "INCOIS ERDDAP",
        "10-day analysis, 1° grid; latest field 30 Jul 2026",
        "Coarse cross-check on SST",
        "10-day; 1°",
        "ERDDAP griddap",
        "Live but stale; low confidence",
    ),
    (
        "Fetch manifest",
        "fetch_manifest.csv",
        "Record of every request that built the snapshot",
        "—",
        "URL, HTTP status, latency, bytes, time for each fetch",
        "Provenance",
        "—",
        "—",
        "—",
    ),
]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

MAX_ROWS = 50000


def add_csv_sheet(wb: Workbook, title: str, path: Path) -> int:
    ws = wb.create_sheet(title[:31])
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        ws.append(header)
        n = 0
        for row in reader:
            out = []
            for v in row:
                try:
                    out.append(
                        float(v)
                        if ("." in v or "e" in v.lower())
                        and v.replace(".", "", 1)
                        .replace("-", "", 1)
                        .replace("e", "", 1)
                        .replace("+", "", 1)
                        .isdigit()
                        else int(v)
                        if v.lstrip("-").isdigit()
                        else v
                    )
                except ValueError:
                    out.append(v)
            ws.append(out)
            n += 1
            if n >= MAX_ROWS:
                break
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 4))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return n


def main() -> None:
    wb = Workbook()
    cat = wb.active
    cat.title = "Catalogue"
    cat.append(
        ["Sagar Setu — datasets in use (all free; extracts of real responses, 15–18 Sep 2026)"]
    )
    cat["A1"].font = Font(name=FONT, bold=True, size=13)
    cat.append([])
    header = [
        "Sheet",
        "Dataset",
        "Publisher",
        "What it contains",
        "Role in Sagar Setu",
        "Cadence / resolution",
        "Access",
        "Status",
        "Rows in sample",
        "Sample file",
    ]
    cat.append(header)
    for c in cat[3]:
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, WRAP
    for sheet, file, name, pub, what, role, cadence, access, status in CATALOGUE:
        p = D / file
        if not p.exists():
            print(f"  missing {file}, skipped")
            continue
        n = add_csv_sheet(wb, sheet, p)
        cat.append([sheet, name, pub, what, role, cadence, access, status, n, file])
    for row in cat.iter_rows(min_row=4):
        for c in row:
            c.font, c.alignment = BODY_FONT, WRAP
    for col, w in zip("ABCDEFGHIJ", (22, 34, 30, 60, 40, 22, 30, 24, 14, 38), strict=False):
        cat.column_dimensions[col].width = w
    cat.freeze_panes = "A4"
    r = cat.max_row + 2
    cat.cell(r, 1, "Notes").font = Font(name=FONT, bold=True)
    notes = [
        "Every value in this workbook was fetched from the named source by scripts/snapshot.py or scripts/export_datasets.py; nothing is synthesised.",
        "Evaluated and not used: INCOIS ERDDAP gridded SST/wind/chlorophyll (stale since 2006–2023); IMD API (requires institutional key — application pending); INCOIS landing-centre advisory fields (frozen at 27 Apr 2024).",
        "Full dataset cards, limits and open questions: docs/datasets.md. Request shapes and schemas: docs/data-sources.md.",
    ]
    for i, t in enumerate(notes, 1):
        c = cat.cell(r + i, 1, t)
        c.font, c.alignment = BODY_FONT, Alignment(wrap_text=False)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1e6:.1f} MB, {len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
